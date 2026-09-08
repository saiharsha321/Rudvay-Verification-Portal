import io
import csv
from typing import List, Dict, Any, Tuple
import openpyxl
from functions.shared.errors import InvalidExcelError
from functions.config import settings

def parse_csv_or_excel(
    file_bytes: bytes,
    filename: str
) -> Tuple[List[str], List[Dict[str, Any]]]:
    """
    Parses a CSV or XLSX file stream and returns (headers, rows).
    Headers are stripped and validated.
    """
    if len(file_bytes) > settings.MAX_EXCEL_FILE_SIZE_BYTES:
        raise InvalidExcelError(f"File size exceeds limit of {settings.MAX_EXCEL_FILE_SIZE_BYTES / (1024*1024)}MB")
    
    lower_name = filename.lower()
    headers: List[str] = []
    rows: List[Dict[str, Any]] = []
    
    if lower_name.endswith('.csv'):
        try:
            content = file_bytes.decode('utf-8-sig', errors='replace')
            reader = csv.reader(io.StringIO(content))
            header_row = next(reader, None)
            if not header_row:
                raise InvalidExcelError("CSV file is empty")
            
            headers = [h.strip() for h in header_row if h.strip()]
            for idx, r in enumerate(reader, start=2):
                if not any(r):  # skip completely empty rows
                    continue
                row_dict = {}
                for col_idx, col_name in enumerate(headers):
                    val = r[col_idx].strip() if col_idx < len(r) else ""
                    row_dict[col_name] = val
                row_dict["_row_number"] = idx
                rows.append(row_dict)
                if len(rows) > settings.MAX_EXCEL_ROWS:
                    raise InvalidExcelError(f"Row count exceeds maximum limit of {settings.MAX_EXCEL_ROWS}")
        except Exception as e:
            if isinstance(e, InvalidExcelError):
                raise e
            raise InvalidExcelError(f"Failed to parse CSV file: {str(e)}")
            
    elif lower_name.endswith('.xlsx') or lower_name.endswith('.xls'):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet = wb.active
            if not sheet:
                raise InvalidExcelError("Excel workbook contains no active sheet")
            
            iter_rows = sheet.iter_rows(values_only=True)
            header_row = next(iter_rows, None)
            if not header_row:
                raise InvalidExcelError("Excel sheet is empty")
            
            headers = [str(h).strip() for h in header_row if h is not None and str(h).strip()]
            if not headers:
                raise InvalidExcelError("No valid header columns detected in Excel sheet")
            
            for idx, r in enumerate(iter_rows, start=2):
                if not r or not any(r):
                    continue
                row_dict = {}
                for col_idx, col_name in enumerate(headers):
                    val = r[col_idx] if col_idx < len(r) else ""
                    val_str = str(val).strip() if val is not None else ""
                    row_dict[col_name] = val_str
                row_dict["_row_number"] = idx
                rows.append(row_dict)
                if len(rows) > settings.MAX_EXCEL_ROWS:
                    raise InvalidExcelError(f"Row count exceeds maximum limit of {settings.MAX_EXCEL_ROWS}")
            wb.close()
        except Exception as e:
            if isinstance(e, InvalidExcelError):
                raise e
            raise InvalidExcelError(f"Failed to parse XLSX file: {str(e)}")
    else:
        raise InvalidExcelError("Unsupported file format. Please upload .xlsx or .csv")
        
    return headers, rows
